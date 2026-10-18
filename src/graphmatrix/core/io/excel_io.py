"""Excel (.xlsx/.xlsm) import/export via openpyxl."""

from __future__ import annotations

from pathlib import Path

from ..model import Workbook, Worksheet
from ._types import build_worksheet


def read_workbook(path: str | Path) -> Workbook:
    from openpyxl import load_workbook

    path = Path(path)
    wb = load_workbook(path, read_only=True, data_only=True)
    book = Workbook(path.stem)
    try:
        for ws in wb.worksheets:
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
            rows = [r for r in rows if any(c is not None and c != "" for c in r)]
            if not rows:
                book.add_sheet(Worksheet(ws.title))
                continue
            header, body = rows[0], rows[1:]
            book.add_sheet(build_worksheet(ws.title, header, body))
    finally:
        wb.close()
    return book


def write_workbook(book: Workbook, path: str | Path) -> Path:
    from openpyxl import Workbook as XLWorkbook

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    xl = XLWorkbook()
    xl.remove(xl.active)
    for sheet in book.sheets:
        ws = xl.create_sheet(title=sheet.name[:31])
        ws.append(
            [f"{c.name} ({c.units})" if c.units else c.name for c in sheet.columns]
        )
        for row in range(sheet.num_rows):
            ws.append(
                [
                    column.get(row) if row < len(column) else None
                    for column in sheet.columns
                ]
            )
    if not book.sheets:
        xl.create_sheet(title="Sheet1")
    xl.save(path)
    return path


__all__ = ["read_workbook", "write_workbook"]
